import io, json, os, shutil, tempfile, unittest, zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from engine import convert, detect, pages, ConversionError, office
from server import unpack

class EngineTests(unittest.TestCase):
    def setUp(self): self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def image(self):
        p=self.root/'source.png'; Image.new('RGBA',(120,80),(255,0,0,0)).save(p); return p
    def doc(self):
        p=self.root/'source.docx'; d=Document(); d.add_paragraph('FileFlip editable text')
        t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Item'; t.cell(0,1).text='Cost'; t.cell(1,0).text='=1+1'; t.cell(1,1).text='12'; d.save(p); return p
    def pdf(self):
        p=self.root/'source.pdf'; c=canvas.Canvas(str(p)); c.drawString(60,750,'FileFlip searchable document'); c.showPage(); c.drawString(60,750,'Second page'); c.save(); return p
    def test_actual_jpeg_encoding_and_alpha_background(self):
        r=convert([self.image()],{'target':'jpg','background':'#00FF00'},self.root)
        f=self.root/r['file']; self.assertEqual(detect(f),'jpg')
        with Image.open(f) as im:
            pixel=im.getpixel((10,10)); self.assertLess(pixel[0],5); self.assertGreater(pixel[1],250)
    def test_jpeg_to_png_encoding(self):
        p=self.root/'in.jpg'; Image.new('RGB',(10,10),'blue').save(p)
        r=convert([p],{'target':'png'},self.root); self.assertEqual(detect(self.root/r['file']),'png')
    def test_rotation(self):
        r=convert([self.image()],{'target':'jpg','rotations':[90]},self.root)
        with Image.open(self.root/r['file']) as im: self.assertEqual(im.size,(80,120))
    def test_docx_tables_are_editable_literal_cells(self):
        r=convert([self.doc()],{'target':'xlsx'},self.root)
        w=load_workbook(self.root/r['file']); self.assertEqual(w.active['A2'].value,'=1+1'); self.assertEqual(w.active['A2'].data_type,'s'); self.assertTrue(r['preview'])
    def test_non_table_document_rejected(self):
        p=self.root/'plain.docx'; d=Document(); d.add_paragraph('Not a table'); d.save(p)
        with self.assertRaisesRegex(ConversionError,'No usable table'): convert([p],{'target':'xlsx'},self.root)
    def test_pdf_editable_word(self):
        r=convert([self.pdf()],{'target':'docx','pages':'2'},self.root)
        text='\n'.join(p.text for p in Document(self.root/r['file']).paragraphs)
        self.assertIn('Second page',text); self.assertNotIn('searchable',text)
    def test_pdf_page_images_zip(self):
        r=convert([self.pdf()],{'target':'png'},self.root)
        with zipfile.ZipFile(self.root/r['file']) as z:
            self.assertEqual(len(z.namelist()),2)
            for n in z.namelist(): self.assertTrue(z.read(n).startswith(b'\x89PNG'))
    def test_pdf_ruled_table(self):
        p=self.root/'table.pdf'; t=Table([['Item','Price'],['Apple','10'],['Pear','20']]); t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)])); SimpleDocTemplate(str(p)).build([t])
        r=convert([p],{'target':'xlsx'},self.root); self.assertEqual(load_workbook(self.root/r['file']).active['A2'].value,'Apple')
    def test_excel_word_range_and_sheet_selection(self):
        p=self.root/'source.xlsx'; w=Workbook(); w.active.title='Keep'; w.active.append(['A','B']); w.active.append([1,2]); w.active.print_area='A1:B2'; w.create_sheet('Drop'); w.save(p)
        r=convert([p],{'target':'docx','sheets':'Keep','orientation':'landscape'},self.root)
        d=Document(self.root/r['file']); self.assertEqual(len(d.tables),1); self.assertEqual(d.tables[0].cell(1,1).text,'2')
    def test_images_embedded_docx(self):
        r=convert([self.image()],{'target':'docx','imageMode':'embed'},self.root); self.assertEqual(len(Document(self.root/r['file']).inline_shapes),1)
    @unittest.skipUnless(shutil.which('tesseract'),'Tesseract unavailable')
    def test_real_image_ocr(self):
        p=self.root/'text.png'; im=Image.new('RGB',(1200,250),'white'); font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',54); ImageDraw.Draw(im).text((30,70),'FileFlip test 123',font=font,fill='black'); im.save(p)
        r=convert([p],{'target':'docx','imageMode':'ocr'},self.root); text=' '.join(p.text for p in Document(self.root/r['file']).paragraphs); self.assertIn('123',text)
    def table_image(self):
        p=self.root/'table-scan.png'; im=Image.new('RGB',(1000,450),'white'); font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40); draw=ImageDraw.Draw(im)
        for y,row in zip((40,150,260), [('Item','Price'),('Apple','10'),('Pear','20')]):
            draw.text((40,y),row[0],font=font,fill='black'); draw.text((520,y),row[1],font=font,fill='black')
        im.save(p); return p
    def test_ocr_image_table(self):
        r=convert([self.table_image()],{'target':'xlsx'},self.root)
        self.assertEqual(load_workbook(self.root/r['file']).active['B2'].value,'10')
    def test_scanned_pdf_ocr(self):
        image=self.table_image(); p=self.root/'scanned.pdf'; c=canvas.Canvas(str(p)); c.drawImage(str(image),30,400,width=540,height=243); c.save()
        r=convert([p],{'target':'docx'},self.root)
        self.assertIn('Apple',' '.join(p.text for p in Document(self.root/r['file']).paragraphs))
    def test_scanned_pdf_table(self):
        image=self.table_image(); p=self.root/'scanned-table.pdf'; c=canvas.Canvas(str(p)); c.drawImage(str(image),30,400,width=540,height=243); c.save()
        r=convert([p],{'target':'xlsx'},self.root)
        self.assertEqual(load_workbook(self.root/r['file']).active['B2'].value,'10')
    def test_invalid_file_rejected(self):
        p=self.root/'fake.pdf'; p.write_text('not a PDF')
        with self.assertRaises(ConversionError): convert([p],{'target':'docx'},self.root)
    def test_page_range_validation(self):
        self.assertEqual(pages('1,3-4,3',4),[0,2,3])
        for s in ('0','3-2','1-9','x','1-2-3'):
            with self.assertRaises(ConversionError): pages(s,4)
    def test_zip_path_traversal_rejected(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,'w') as z:
            z.writestr('manifest.json',json.dumps({'files':['../escape.pdf'],'options':{}})); z.writestr('../escape.pdf','bad')
        with self.assertRaises(ConversionError): unpack(data.getvalue(),self.root)
        self.assertFalse((self.root.parent/'escape.pdf').exists())
    def test_legacy_output_rejected(self):
        with self.assertRaises(ConversionError): convert([self.image()],{'target':'doc'},self.root)
    @unittest.skipUnless(shutil.which(os.environ.get('FILEFLIP_SOFFICE','libreoffice')),'LibreOffice unavailable')
    def test_office_pdf_real_render(self):
        r=convert([self.doc()],{'target':'pdf'},self.root); self.assertEqual(detect(self.root/r['file']),'pdf')
    @unittest.skipUnless(shutil.which(os.environ.get('FILEFLIP_SOFFICE','libreoffice')),'LibreOffice unavailable')
    def test_excel_pdf_real_render(self):
        p=self.root/'book.xlsx'; w=Workbook(); w.active.append(['Header','Value']); w.active.append(['Hello',10]); w.save(p)
        r=convert([p],{'target':'pdf','orientation':'landscape','range':'A1:B2'},self.root); self.assertEqual(detect(self.root/r['file']),'pdf')
    @unittest.skipUnless(shutil.which(os.environ.get('FILEFLIP_SOFFICE','libreoffice')) and __import__('importlib').util.find_spec('olefile'),'LibreOffice/olefile unavailable')
    def test_legacy_doc_rendering(self):
        legacy=office(self.doc(),'doc',self.root)
        self.assertEqual(detect(legacy),'doc')
        r=convert([legacy],{'target':'pdf'},self.root)
        self.assertEqual(detect(self.root/r['file']),'pdf')
    @unittest.skipUnless(shutil.which(os.environ.get('FILEFLIP_SOFFICE','libreoffice')) and __import__('importlib').util.find_spec('olefile'),'LibreOffice/olefile unavailable')
    def test_legacy_xls_input(self):
        p=self.root/'legacy-source.xlsx'; w=Workbook(); w.active.append(['Name','Score']); w.active.append(['Ada',99]); w.save(p)
        legacy=office(p,'xls',self.root); self.assertEqual(detect(legacy),'xls')
        r=convert([legacy],{'target':'docx'},self.root)
        self.assertEqual(Document(self.root/r['file']).tables[0].cell(1,1).text,'99')
if __name__=='__main__': unittest.main()
